"""Builds the Praedium Group Agent Score Card workbook (.xlsx).

The file is uploaded to Google Drive and converted to a Google Sheet.
Automation (weekly emails, sharing, tab protection) lives in Code.gs.

Run:  python3 scorecard/build_scorecard.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NUM_AGENTS = 7
OWNER_EMAIL = "ncarnes@kw.com"
# (name, email) — empty slots stay as "Agent N" until filled in on Settings
AGENTS = [
    ("Gabrielle Gillie", "gabeg@kwcommercial.com"),
    ("Dillon", "dillon2403@gmail.com"),
    ("Marcus Lominick", "mlominick@kwcommercial.com"),
    ("Mary Elizabeth Young", "myoung@kwcommercial.com"),
    ("Jayan Abraham", "jayanabraham@kw.com"),
]
FIRST_WEEK = dt.date(2026, 9, 28)  # Monday
LAST_WEEK = dt.date(2027, 12, 27)
Q4_START, Q4_END = dt.date(2026, 10, 1), dt.date(2026, 12, 31)

NAVY, GOLD, LIGHT, INPUT, GREY = "1F3A5F", "C9A227", "EEF2F7", "FFFBEA", "F3F4F6"
F_TITLE = Font(name="Arial", size=16, bold=True, color="FFFFFF")
F_H = Font(name="Arial", size=11, bold=True, color="FFFFFF")
F_SEC = Font(name="Arial", size=12, bold=True, color=NAVY)
F_B = Font(name="Arial", size=10)
F_BB = Font(name="Arial", size=10, bold=True)
F_NOTE = Font(name="Arial", size=9, italic=True, color="6B7280")
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_GOLD = PatternFill("solid", fgColor=GOLD)
FILL_LIGHT = PatternFill("solid", fgColor=LIGHT)
FILL_INPUT = PatternFill("solid", fgColor=INPUT)
FILL_GREY = PatternFill("solid", fgColor=GREY)
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

MONEY = '"$"#,##0'
PCT = "0%"
DATE = "mmm d, yyyy"

SETTINGS = "Settings"
AGENT_ROW0 = 14  # first agent row on Settings


def agent_label(i):
    return AGENTS[i - 1][0] if i <= len(AGENTS) else f"Agent {i}"


def agent_email(i):
    return AGENTS[i - 1][1] if i <= len(AGENTS) else ""


def scorecard_name(i):
    return f"{agent_label(i)} Scorecard"


def assessment_name(i):
    return f"{agent_label(i)} Assessment"


def q(name):
    return "'" + name.replace("'", "''") + "'"


def title(ws, text, width_cols, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    c = ws.cell(1, 1, text)
    c.font, c.fill, c.alignment = F_TITLE, FILL_NAVY, Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
        s = ws.cell(2, 1, subtitle)
        s.font, s.alignment = F_NOTE, Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[2].height = 30


def header_row(ws, row, labels, start_col=1, fill=FILL_NAVY):
    for j, lab in enumerate(labels):
        c = ws.cell(row, start_col + j, lab)
        c.font, c.fill, c.alignment, c.border = F_H, fill, CENTER, BOX


def list_dv(ws, options, sqref):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.add(sqref)
    ws.add_data_validation(dv)


# --------------------------------------------------------------------------- #
# Assessment content (from the rough draft)
# --------------------------------------------------------------------------- #
YN = ["Yes", "No"]
TYPE_OF_BIZ = ["Primarily Commercial", "Primarily Residential", "Mix of Commercial and Residential",
               "New to Real Estate", "Building My Commercial Business", "Other"]
SOURCES = ["Cold Calling Property Owners", "Calling Businesses / Tenants", "Sphere of Influence",
           "Past Clients", "Current Clients", "KW Residential Agent Referrals", "KW Commercial Referrals",
           "Outside Broker Referrals", "Attorneys", "CPAs / Accountants", "Lenders / Bankers",
           "Developers / Builders", "Investors", "Business Owners", "Networking Organizations",
           "Email Campaigns", "Direct Mail", "Social Media / Content", "Crexi / Online Platforms",
           "Database / CRM Follow-Up", "Property Visits / Canvassing", "Market Center Relationships", "Other"]
SKILLS = ["Cold Calling", "Prospecting", "Building Prospect Lists", "Getting Decision Makers on the Phone",
          "Converting Conversations into Meetings", "Discovery Meetings", "Listing Presentations",
          "Buyer / Tenant Presentations", "Market Research", "Land id", "Crexi", "GIS",
          "CRM / Database Management", "Financial Underwriting", "NOI / Cap Rate / Valuation",
          "BOV Preparation", "LOIs", "Purchase Agreements", "Commercial Leases", "Due Diligence",
          "Negotiation", "Deal Management", "Client Follow-Up", "Networking", "Referral Development",
          "AI / ChatGPT", "Personal Organization", "Time Management"]
TIME_CATS = ["Commercial Prospecting / Business Development", "Commercial Client & Deal Work",
             "Commercial Market Research", "Follow-Up / CRM", "Networking / Sphere Development",
             "Marketing / Content", "Residential Business", "Administrative Work",
             "Training / Professional Development", "Other"]
PIPELINE = ["Active Listings", "Buyer Assignments", "Tenant Representation", "Landlord Representation",
            "Active LOIs", "Under Contract", "Qualified Prospects", "Early-Stage Opportunities"]
SUPPORT = ["Prospecting", "Prospect Lists", "Market Research", "Deal Strategy", "Underwriting",
           "Contracts / LOIs", "Negotiation", "Listing Presentations", "Marketing", "Technology", "CRM",
           "Administrative Support", "Joint Meetings / Calls", "Mentoring", "Accountability", "Referrals",
           "Market Center Development", "Specialized CRE Training"]
PAIN = ["Not enough prospecting", "Difficulty finding prospects", "Difficulty making calls",
        "Difficulty converting calls", "Follow-up", "Time management", "Daily structure",
        "Residential business consuming time", "Lack of experience", "Lack of confidence",
        "Market knowledge", "Product knowledge", "Financial analysis", "Contracts / LOIs", "Negotiation",
        "Technology", "Marketing", "Building a sphere", "Market Center relationships",
        "Lack of specialization", "Organization / CRM", "Accountability", "Other"]
MC_ACTIONS = ["Introduced myself individually to agents", "Attended team/office meetings",
              "Given commercial real estate presentations", "Shared commercial opportunities/listings",
              "Asked agents directly for referrals", "Helped agents with commercial questions",
              "Developed relationships with leadership", "Participated in office events",
              "Sent emails/newsletters to agents", "Shared CRE educational content",
              "Nothing consistently yet", "Other"]
SPHERE_WHO = ["Friends / Family", "Former Clients", "Residential Clients", "Business Owners", "Investors",
              "Attorneys", "CPAs", "Lenders", "Developers", "Contractors", "Other Brokers",
              "Community Organizations", "Other"]
PRODUCTS = ["Retail", "Office", "Industrial", "Land", "Multifamily", "Medical / Healthcare", "Hospitality",
            "Investment Sales", "Development", "Business Brokerage", "Other"]
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
TIME_GUIDE = [
    ("Prospecting / Business Development", "Cold calling, owner outreach, tenant outreach, email campaigns, canvassing, prospect research"),
    ("Client / Deal Work", "Tours, client meetings, LOIs, contracts, negotiations, due diligence, inspections, closing coordination"),
    ("Market Research", "Crexi, CoStar, GIS, Land id, researching ownership, comps, zoning, demographics"),
    ("Follow-Up / CRM", "Returning calls, email follow-up, CRM updates, maintaining prospect lists"),
    ("Networking / Sphere", "KW agents, attorneys, lenders, CPAs, developers, business owners, networking events"),
    ("Marketing", "Social media, property marketing, flyers, listing materials, content creation"),
    ("Residential", "Residential prospecting, showings, listings, buyers, contracts, closings"),
    ("Administrative", "Paperwork, scheduling, expense tracking, file organization, general admin"),
    ("Training / Development", "Courses, meetings, CRE training, technology training, role-playing"),
    ("Other", "Anything that does not reasonably fit above"),
]


class Assessment:
    """Writes the assessment one row at a time.  Column F flags rows that
    count toward the completion %; answers always start in column C."""

    def __init__(self, ws):
        self.ws, self.r, self.keys = ws, 4, {}

    def _label(self, text, bold=False):
        c = self.ws.cell(self.r, 2, text)
        c.font, c.alignment, c.border = (F_BB if bold else F_B), WRAP, BOX

    def _answer(self, fmt=None, merge=True, long=False, flag=True):
        ws, r = self.ws, self.r
        if merge:
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        for col in (3, 4, 5) if merge else (3,):
            ws.cell(r, col).fill, ws.cell(r, col).border = FILL_INPUT, BOX
        a = ws.cell(r, 3)
        a.alignment, a.font = WRAP, F_B
        if fmt:
            a.number_format = fmt
        if long:
            ws.row_dimensions[r].height = 48
        if flag:
            ws.cell(r, 6, 1)
        return a

    def section(self, text, note=None):
        self.r += 1
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=5)
        c = ws.cell(self.r, 1, text)
        c.font, c.fill, c.alignment = F_H, FILL_NAVY, Alignment(vertical="center", indent=1)
        ws.row_dimensions[self.r].height = 24
        self.r += 1
        if note:
            ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=5)
            n = ws.cell(self.r, 1, note)
            n.font, n.alignment = F_NOTE, Alignment(wrap_text=True, indent=1)
            ws.row_dimensions[self.r].height = 28
            self.r += 1

    def q(self, text, kind="text", options=None, key=None):
        self._label(text)
        fmt = {"money": MONEY, "pct": PCT}.get(kind)
        a = self._answer(fmt=fmt, long=(kind == "long"))
        if options:
            list_dv(self.ws, options, a.coordinate)
        if key:
            self.keys[key] = a.coordinate
        self.r += 1

    def multi(self, text, options):
        """Check-all-that-apply: one Yes/No dropdown per option."""
        self._label(text + "  (select Yes for all that apply)", bold=True)
        self.r += 1
        start = self.r
        for opt in options:
            c = self.ws.cell(self.r, 2, "      " + opt)
            c.font, c.border = F_B, BOX
            a = self.ws.cell(self.r, 3)
            a.fill, a.border, a.alignment = FILL_INPUT, BOX, CENTER
            if opt == "Other":
                self.ws.merge_cells(start_row=self.r, start_column=4, end_row=self.r, end_column=5)
                d = self.ws.cell(self.r, 4)
                d.fill, d.border = FILL_INPUT, BOX
                self.ws.cell(self.r, 2).value = "      Other (describe →)"
            self.r += 1
        list_dv(self.ws, YN, f"C{start}:C{self.r - 1}")

    def grid(self, cols, rows, validations=None, fmts=None, total=False, pct_of_col=None):
        """Table: B = row label, answers in C.. (up to E). Only column C counts
        toward completion."""
        header_row(self.ws, self.r, cols, start_col=2, fill=FILL_GOLD)
        self.r += 1
        start = self.r
        n = len(cols) - 1
        for lab in rows:
            c = self.ws.cell(self.r, 2, lab)
            c.font, c.border = F_B, BOX
            for j in range(n):
                a = self.ws.cell(self.r, 3 + j)
                a.fill, a.border, a.alignment = FILL_INPUT, BOX, CENTER
                if fmts and fmts[j]:
                    a.number_format = fmts[j]
            self.ws.cell(self.r, 6, 1)
            self.r += 1
        end = self.r - 1
        for j, opts in enumerate(validations or []):
            if opts:
                col = get_column_letter(3 + j)
                list_dv(self.ws, opts, f"{col}{start}:{col}{end}")
        if pct_of_col is not None:
            # column D = share of the column C total
            for rr in range(start, end + 1):
                d = self.ws.cell(rr, 4, f'=IFERROR(C{rr}/SUM(C${start}:C${end}),"")')
                d.fill, d.number_format = FILL_GREY, PCT
        if total:
            c = self.ws.cell(self.r, 2, "TOTAL")
            c.font, c.fill, c.border = F_BB, FILL_LIGHT, BOX
            for j in range(n):
                col = get_column_letter(3 + j)
                t = self.ws.cell(self.r, 3 + j, f"=SUM({col}{start}:{col}{end})")
                t.font, t.fill, t.border, t.alignment = F_BB, FILL_LIGHT, BOX, CENTER
                if fmts and fmts[j]:
                    t.number_format = fmts[j]
            self.r += 1
        return start, end

    def guide(self, rows):
        header_row(self.ws, self.r, ["Category", "Examples"], start_col=2, fill=FILL_GREY)
        for c in (2, 3):
            self.ws.cell(self.r, c).font = F_BB
        self.r += 1
        for cat, ex in rows:
            self.ws.cell(self.r, 2, cat).font = F_BB
            self.ws.merge_cells(start_row=self.r, start_column=3, end_row=self.r, end_column=5)
            e = self.ws.cell(self.r, 3, ex)
            e.font, e.alignment = F_NOTE, WRAP
            self.ws.row_dimensions[self.r].height = 26
            self.r += 1

    def gap(self):
        self.r += 1


def build_assessment(ws, i):
    name_ref = f"{q(SETTINGS)}!B{AGENT_ROW0 + i - 1}"
    ws.merge_cells("A1:E1")
    t = ws.cell(1, 1, f'=IF({name_ref}="","Agent {i}",{name_ref})&" — Business Development & Growth Assessment"')
    t.font, t.fill, t.alignment = F_TITLE, FILL_NAVY, Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:E2")
    s = ws.cell(2, 1, "THE PRAEDIUM GROUP  •  This assessment helps us understand where you are today, where your "
                      "time is going, and what training, structure, or support could help you grow. There are no "
                      "“right” answers — answer based on what your business actually looks like today. "
                      "Estimates are completely acceptable. Yellow cells are yours to fill in.")
    s.font, s.alignment = F_NOTE, Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[2].height = 44
    ws.cell(3, 2, "Completion:").font = F_BB
    ws.cell(3, 2).alignment = Alignment(horizontal="right")
    pc = ws.cell(3, 3, '=IFERROR(SUMPRODUCT((F5:F400=1)*(C5:C400<>""))/SUM(F5:F400),0)')
    pc.number_format, pc.font, pc.fill = PCT, Font(name="Arial", size=12, bold=True, color=NAVY), FILL_LIGHT

    A = Assessment(ws)
    A.section("1. ABOUT YOUR BUSINESS TODAY")
    A.q("Market Center / KW Office")
    A.q("How long have you been licensed in real estate?")
    A.q("How long have you actively worked in commercial real estate?")
    A.q("How would you describe your current real estate business?", options=TYPE_OF_BIZ)
    A.q("If you selected Other, please explain:", "long")
    A.q("Is real estate currently your primary occupation?", options=["Yes", "No", "Transitioning toward it"])
    A.q("If not, what other work or major commitments currently compete for your time?", "long")

    A.section("2. YOUR CURRENT WORK WEEK", "We want to understand what your schedule actually looks like today.")
    A.q("On average, how many days per week do you currently spend working on real estate?")
    A.q("How many hours do you work on real estate during a typical day?")
    A.q("What time do you normally begin your workday?")
    A.q("What time do you normally finish your workday?")
    A.q("Is your schedule generally consistent from day to day?", options=["Yes", "Somewhat", "No"])
    A.gap()
    A.grid(["Day", "Usually Work?", "Approx. Hours"], DAYS, validations=[YN, None])
    A.gap()
    A._label("How is your typical real estate week divided?", bold=True)
    A.r += 1
    A.grid(["Area", "Hours / Week"], ["Commercial Real Estate", "Residential Real Estate",
                                      "Administrative / Business Operations", "Other"], total=True)
    A.q("If Other, what does that time consist of?", "long")

    A.section("3. WHERE DOES YOUR TIME ACTUALLY GO?",
              "You do not need to track every minute. Think about a typical week and make your best estimate. "
              "Use this guide if you are unsure how to categorize something.")
    A.guide(TIME_GUIDE)
    A.gap()
    s3, e3 = A.grid(["Activity", "Approx. Hours / Week", "% of Real Estate Time"], TIME_CATS,
                    fmts=[None, PCT], total=True, pct_of_col=True)
    ws.cell(A.r - 1, 4).value = f"=IFERROR(SUM(D{s3}:D{e3}),\"\")"
    ws.cell(A.r - 1, 4).number_format = PCT
    A.q("If you listed time under Other, please explain:", "long")
    A.q("Which activities take up the MOST time in your week?", "long")
    A.q("Which activities are currently producing the MOST business for you?", "long")
    A.q("What do you feel you are spending too much time doing?", "long")
    A.q("What do you feel you are not spending enough time doing?", "long")

    A.section("4. HOW DO YOU STRUCTURE YOUR DAY?")
    A.q("Do you currently plan your workday in advance?", options=["Almost always", "Sometimes", "Rarely", "Never"])
    A.q("Do you currently use time blocks for specific activities?", options=["Yes", "Somewhat", "No"])
    A.q("Do you have a dedicated prospecting period during your day?", options=["Yes", "Sometimes", "No"])
    A.q("If yes — typical prospecting time (e.g. 9:00–11:00 AM):")
    A.q("If yes — typical length:")
    A.q("When do you feel you are most productive during the day?", "long")
    A.q("What normally causes your schedule to get off track?", "long")
    A.q("What is the first business-producing activity you normally do each day?", "long")
    A.q("What usually consumes more time than you intended?", "long")

    A.section("5. HOW ARE YOU CURRENTLY GENERATING BUSINESS?",
              "Activity level: 0 = Not currently using · 1 = Occasionally · 3 = Consistently · 5 = Major part of my business")
    A.grid(["Business Development Source", "Activity Level 0–5", "Generated a Real Opportunity?"], SOURCES,
           validations=[["0", "1", "2", "3", "4", "5"], YN])
    A.q("If Other, please explain:", "long")
    A.q("What are currently your TOP THREE sources of business?", "long")
    A.q("Which one has produced the most actual opportunities or transactions?", "long")
    A.q("Which business-development channel do you believe you should be using more?", "long")

    A.section("6. YOUR MARKET CENTER")
    A.q("How have you been developing relationships within your Market Center?", "long")
    A.q("How many residential agents in your Market Center know that you handle commercial real estate?",
        options=["Very few", "Some", "Most", "Nearly all", "Unsure"])
    A.multi("What have you done to generate commercial referrals from your Market Center?", MC_ACTIONS)
    A.q("Meaningful conversations with Market Center agents in a typical month:")
    A.q("Commercial referrals received from your Market Center in the past 90 days:")
    A.q("What could you do differently over the next 90 days to become the commercial resource within your office?", "long")

    A.section("7. YOUR SPHERE OF INFLUENCE")
    A.multi("Who currently knows that you work in commercial real estate?", SPHERE_WHO)
    A.q("Approximately how many people are part of your professional sphere today?")
    A.q("How often are you intentionally communicating with your sphere?",
        options=["Weekly", "Monthly", "Quarterly", "Occasionally", "Almost never"])
    A.q("What are you currently doing to grow your sphere?", "long")
    A.q("Five people or categories of people you should be building stronger relationships with:", "long")

    A.section("8. MARKET & PRODUCT FOCUS")
    A.q("Do you currently have a defined commercial real estate focus?",
        options=["Yes", "Somewhat", "No", "Still determining it"])
    A.q("Are you focusing on a particular geographic area?", options=YN)
    A.q("Primary Market:")
    A.q("Secondary Market:")
    A.q("Specific submarkets / corridors / neighborhoods:", "long")
    A.multi("Which property types are you currently pursuing?", PRODUCTS)
    A.q("Which ONE or TWO areas would you most like to become known for?", "long")
    A.q("If someone asked you today, “What type of commercial real estate do you specialize in?” what would you say?", "long")
    A.q("Do you feel confident saying that? Why or why not?", "long")

    A.section("9. CURRENT COMMERCIAL PIPELINE", "Do your best. Exact numbers are not required.")
    A.grid(["Category", "#", "Approx. Transaction Volume", "Potential GCI"], PIPELINE,
           fmts=[None, MONEY, MONEY], total=True)
    A.keys["pipeline_gci"] = f"E{A.r - 1}"
    A.q("New commercial opportunities created in the last 30 days:")
    A.q("Prospect/client meetings in the last 30 days:")
    A.q("New people added to your commercial database in the last 30 days:")
    A.q("How many follow-ups are currently outstanding? (number or “Unsure”)")
    A.q("Do you currently know what your commercial pipeline is worth in potential GCI?",
        options=["Yes", "Approximately", "No"])

    A.section("10. SKILLS & CONFIDENCE ASSESSMENT",
              "Rate yourself honestly.  1 = Need substantial help · 3 = Comfortable · 5 = Strong enough to teach someone else")
    s10, e10 = A.grid(["Area", "Score 1–5"], SKILLS, validations=[["1", "2", "3", "4", "5"]])
    A.keys["skills_range"] = f"C{s10}:C{e10}"
    A.q("Three areas where you feel strongest:", "long")
    A.q("Three areas where you most want help:", "long")

    A.section("11. PAIN POINTS & AREAS FOR IMPROVEMENT")
    A.multi("What do you believe is currently preventing you from growing faster?", PAIN)
    A.q("What is your single biggest pain point right now?", "long")
    A.q("What is one recurring problem in your business that you would like help solving?", "long")
    A.q("What is something you know you should be doing but have not consistently done?", "long")

    A.section("12. PRAEDIUM SUPPORT & DEVELOPMENT")
    A.grid(["Area", "Need More Support?"], SUPPORT, validations=[YN])
    A.q("What could Praedium do that would make the biggest difference in your business?", "long")
    A.q("What do you need more of from me personally?", "long")
    A.q("What should we be doing as a team that we are currently not doing?", "long")
    A.q("What would you most like to learn during our biweekly meetings?", "long")

    A.section("13. Q4 GOALS  (Oct 1 – Dec 31)",
              "These goals feed your Weekly Scorecard — progress is tracked against them automatically.")
    A.q("Commercial Transaction Volume Goal ($)", "money", key="q4_volume")
    A.q("Commercial GCI Goal ($)", "money", key="q4_gci")
    A.q("Number of Closings", key="q4_closings")
    A.q("New Listings", key="q4_listings")
    A.q("New Buyer / Tenant Assignments", key="q4_assignments")
    A.q("Qualified Opportunities Created", key="q4_opps")
    A.q("Meetings Per Month", key="q4_meetings_pm")
    A.q("The three most important things you want to accomplish before December 31:", "long")

    A.section("14. 2027 GOALS")
    A.q("2027 Commercial Volume Goal ($)", "money", key="y27_volume")
    A.q("2027 Commercial GCI Goal ($)", "money", key="y27_gci")
    A.q("Transactions Closed", key="y27_closings")
    A.q("Listings Obtained", key="y27_listings")
    A.q("Primary Geography")
    A.q("Primary Product Type")
    A.q("What % of your business would you LIKE to be commercial by the end of 2027?", "pct")
    A.q("What would a successful 2027 look like to you?", "long")
    A.q("What would need to change between now and then for that to happen?", "long")

    A.section("15. YOUR NEXT 90 DAYS")
    A.q("One thing I need to STOP doing:", "long")
    A.q("One thing I need to START doing:", "long")
    A.q("One thing I need to CONTINUE doing:", "long")
    A.q("My primary prospecting strategy will be:", "long")
    A.q("My primary market / product focus will be:", "long")
    A.q("The biggest improvement I want to make over the next 90 days is:", "long")
    A.q("One commitment I am making to myself before our next meeting:", "long", key="commitment")

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 62
    for col in "CDE":
        ws.column_dimensions[col].width = 22
    ws.column_dimensions["F"].hidden = True
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = GOLD
    return A.keys


# --------------------------------------------------------------------------- #
# Weekly scorecard
# --------------------------------------------------------------------------- #
WEEKLY_COLS = [
    ("Week Of (Mon)", 13, DATE),
    ("Days Worked", 9, None),
    ("Total RE Hours", 9, None),
    ("Commercial Hours", 10, None),
    ("Prospecting Hours", 11, None),
    ("Calls / Outreach Attempts", 11, None),
    ("Decision-Maker Conversations", 13, None),
    ("Prospect / Client Meetings", 11, None),
    ("New Opportunities Created", 12, None),
    ("New Database Contacts", 11, None),
    ("Follow-Ups Outstanding", 11, None),
    ("Market Center Agent Conversations", 13, None),
    ("Sphere Touches", 10, None),
    ("New Listings", 9, None),
    ("New Buyer / Tenant Assignments", 12, None),
    ("LOIs Submitted", 9, None),
    ("Went Under Contract", 10, None),
    ("Closings", 9, None),
    ("Volume Closed ($)", 13, MONEY),
    ("GCI Closed ($)", 12, MONEY),
    ("Pipeline Potential GCI ($)", 14, MONEY),
    ("Win of the Week", 30, None),
    ("Biggest Obstacle", 30, None),
    ("#1 Focus Next Week", 30, None),
    ("Last Updated (auto)", 17, "mmm d, h:mm am/pm"),
    ("Status", 12, None),
]
COL = {name: get_column_letter(i + 1) for i, (name, _, _) in enumerate(WEEKLY_COLS)}
HEADER_ROW = 14
FIRST_DATA = HEADER_ROW + 1


def weeks():
    d, out = FIRST_WEEK, []
    while d <= LAST_WEEK:
        out.append(d)
        d += dt.timedelta(days=7)
    return out


def build_scorecard(ws, i, akeys):
    n_weeks = len(weeks())
    last = FIRST_DATA + n_weeks - 1
    ncols = len(WEEKLY_COLS)
    aname = q(assessment_name(i))
    name_ref = f"{q(SETTINGS)}!B{AGENT_ROW0 + i - 1}"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(1, 1, f'=IF({name_ref}="","Agent {i}",{name_ref})&" — Weekly Score Card"')
    t.font, t.fill, t.alignment = F_TITLE, FILL_NAVY, Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(2, 1, "Fill in your row for the current week by Friday end of day (yellow cells). Estimates are fine. "
                      "“Follow-Ups Outstanding” and “Pipeline Potential GCI” are snapshots as of the end of the week; "
                      "everything else is what happened that week. Q4 progress below updates automatically.")
    s.font, s.alignment = F_NOTE, Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[2].height = 30

    # Q4 goals vs actual
    ws.merge_cells("A4:F4")
    h = ws.cell(4, 1, "Q4 2026 PROGRESS  (goals come from Section 13 of your Assessment)")
    h.font, h.fill, h.alignment = F_H, FILL_GOLD, Alignment(indent=1, vertical="center")
    header_row(ws, 5, ["Metric", "", "Q4 Goal", "Q4 Actual", "% to Goal", "On Pace?"], fill=FILL_NAVY)
    ws.merge_cells("A5:B5")
    date_rng = f"$A${FIRST_DATA}:$A${last}"
    q4 = f'{date_rng},">="&DATE(2026,9,28),{date_rng},"<="&DATE(2026,12,31)'
    q4_elapsed = "MAX(0,MIN(1,(TODAY()-DATE(2026,10,1))/(DATE(2026,12,31)-DATE(2026,10,1))))"
    goals = [
        ("Commercial GCI", "q4_gci", "GCI Closed ($)", MONEY),
        ("Transaction Volume", "q4_volume", "Volume Closed ($)", MONEY),
        ("Closings", "q4_closings", "Closings", None),
        ("New Listings", "q4_listings", "New Listings", None),
        ("New Buyer / Tenant Assignments", "q4_assignments", "New Buyer / Tenant Assignments", None),
        ("Qualified Opportunities Created", "q4_opps", "New Opportunities Created", None),
        ("Meetings (goal = per month × 3)", "q4_meetings_pm", "Prospect / Client Meetings", None),
    ]
    for k, (label, key, colname, fmt) in enumerate(goals):
        r = 6 + k
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(r, 1, label).font = F_BB
        goal = f"{aname}!{akeys[key]}"
        if key == "q4_meetings_pm":
            goal_f = f'=IF({goal}="","",{goal}*3)'
        else:
            goal_f = f'=IF({goal}="","",{goal})'
        c = COL[colname]
        ws.cell(r, 3, goal_f)
        ws.cell(r, 4, f"=SUMIFS(${c}${FIRST_DATA}:${c}${last},{q4})")
        ws.cell(r, 5, f'=IFERROR(D{r}/C{r},"")')
        ws.cell(r, 6, f'=IF(C{r}="","Set goal",IF(D{r}>=C{r}*{q4_elapsed},"On pace","Behind"))')
        for col in range(1, 7):
            cell = ws.cell(r, col)
            cell.border, cell.alignment = BOX, (CENTER if col > 2 else Alignment(indent=1))
            if col in (3, 4) and fmt:
                cell.number_format = fmt
        ws.cell(r, 5).number_format = PCT
    ws.conditional_formatting.add("F6:F12", CellIsRule(operator="equal", formula=['"On pace"'],
                                  fill=PatternFill("solid", fgColor="D1FAE5")))
    ws.conditional_formatting.add("F6:F12", CellIsRule(operator="equal", formula=['"Behind"'],
                                  fill=PatternFill("solid", fgColor="FEE2E2")))

    # Consistency block
    ws.cell(4, 8, "ACCOUNTABILITY").font = F_SEC
    stats = [
        ("Weeks updated", f'=COUNTIFS({date_rng},"<="&TODAY(),${COL["Last Updated (auto)"]}${FIRST_DATA}:${COL["Last Updated (auto)"]}${last},"<>")'),
        ("Weeks elapsed", f'=COUNTIFS({date_rng},"<="&TODAY())'),
        ("Consistency", "=IFERROR(K5/K6,0)"),
        ("Avg calls / week", f'=IFERROR(AVERAGEIFS(${COL["Calls / Outreach Attempts"]}${FIRST_DATA}:${COL["Calls / Outreach Attempts"]}${last},{date_rng},"<="&TODAY(),${COL["Calls / Outreach Attempts"]}${FIRST_DATA}:${COL["Calls / Outreach Attempts"]}${last},"<>"),0)'),
        ("Avg meetings / week", f'=IFERROR(AVERAGEIFS(${COL["Prospect / Client Meetings"]}${FIRST_DATA}:${COL["Prospect / Client Meetings"]}${last},{date_rng},"<="&TODAY(),${COL["Prospect / Client Meetings"]}${FIRST_DATA}:${COL["Prospect / Client Meetings"]}${last},"<>"),0)'),
        ("Latest pipeline GCI", f'=IFERROR(INDEX(${COL["Pipeline Potential GCI ($)"]}${FIRST_DATA}:${COL["Pipeline Potential GCI ($)"]}${last},MATCH(1E+100,${COL["Pipeline Potential GCI ($)"]}${FIRST_DATA}:${COL["Pipeline Potential GCI ($)"]}${last})),"")'),
        ("Assessment complete", f"={aname}!C3"),
    ]
    for k, (label, f) in enumerate(stats):
        r = 5 + k
        ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
        ws.cell(r, 8, label).font = F_BB
        v = ws.cell(r, 11, f)
        v.border, v.alignment, v.fill = BOX, CENTER, FILL_LIGHT
        v.number_format = {2: PCT, 5: MONEY, 6: PCT}.get(k, "0.0" if k in (3, 4) else "0")

    # Weekly table
    header_row(ws, HEADER_ROW, [c[0] for c in WEEKLY_COLS])
    ws.row_dimensions[HEADER_ROW].height = 48
    for j, (_, width, _) in enumerate(WEEKLY_COLS):
        ws.column_dimensions[get_column_letter(j + 1)].width = width
    upd, status = COL["Last Updated (auto)"], COL["Status"]
    for k, d in enumerate(weeks()):
        r = FIRST_DATA + k
        a = ws.cell(r, 1, d)
        a.number_format, a.font, a.fill, a.border = DATE, F_BB, FILL_LIGHT, BOX
        for j, (_, _, fmt) in enumerate(WEEKLY_COLS[1:-2], start=2):
            c = ws.cell(r, j)
            c.fill, c.border = FILL_INPUT, BOX
            c.alignment = WRAP if j >= 22 else CENTER
            if fmt:
                c.number_format = fmt
        u = ws.cell(r, ncols - 1)
        u.fill, u.border, u.font, u.number_format = FILL_GREY, BOX, F_NOTE, "mmm d, h:mm am/pm"
        st = ws.cell(r, ncols, f'=IF(A{r}>TODAY(),"Upcoming",IF({upd}{r}<>"","✓ Updated",'
                               f'IF(A{r}+6>=TODAY(),"Due Friday","⚠ Missing")))')
        st.border, st.alignment = BOX, CENTER
    rng = f"{status}{FIRST_DATA}:{status}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"⚠ Missing"'],
                                  fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color="991B1B", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"✓ Updated"'],
                                  fill=PatternFill("solid", fgColor="D1FAE5"), font=Font(color="065F46")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Due Friday"'],
                                  fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(color="92400E", bold=True)))
    # highlight the current week's row
    ws.conditional_formatting.add(
        f"A{FIRST_DATA}:{get_column_letter(ncols - 2)}{last}",
        FormulaRule(formula=[f"AND($A{FIRST_DATA}<=TODAY(),$A{FIRST_DATA}+6>=TODAY())"],
                    border=Border(top=Side(style="medium", color=GOLD), bottom=Side(style="medium", color=GOLD))))
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="Numbers only",
                        error="Please enter a number (0 or more).")
    dv.add(f"B{FIRST_DATA}:{COL['Pipeline Potential GCI ($)']}{last}")
    ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(FIRST_DATA, 2)
    ws.sheet_properties.tabColor = NAVY


# --------------------------------------------------------------------------- #
# Settings / Dashboard / Review / Start Here
# --------------------------------------------------------------------------- #
def build_settings(ws):
    title(ws, "Settings — Score Card Automation", 6,
          "Owner-only. Fill in agent names + emails, then use the  Score Card  menu → “Rename tabs from Settings” "
          "and “Share & protect”. Email schedule changes take effect after “Install weekly email triggers”.")
    rows = [
        ("Owner name", "Nicholas Carnes"),
        ("Owner email (weekly team summary goes here)", OWNER_EMAIL),
        ("Team name", "The Praedium Group"),
        ("Agent reminder day", "FRIDAY"),
        ("Agent reminder hour (0–23)", 9),
        ("Owner summary day", "MONDAY"),
        ("Owner summary hour (0–23)", 7),
        ("Send Monday nudge to agents who missed last week?", "Yes"),
    ]
    for k, (lab, val) in enumerate(rows):
        r = 4 + k
        ws.cell(r, 1, lab).font = F_BB
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 3, val)
        c.fill, c.border, c.font = FILL_INPUT, BOX, F_B
    days = ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY"]
    list_dv(ws, days, "C7")
    list_dv(ws, days, "C9")
    list_dv(ws, YN, "C11")

    header_row(ws, AGENT_ROW0 - 1, ["#", "Agent Name", "Agent Email", "Scorecard Tab", "Assessment Tab",
                                    "Market Center"])
    for i in range(1, NUM_AGENTS + 1):
        r = AGENT_ROW0 + i - 1
        vals = [i, agent_label(i) if i <= len(AGENTS) else "", agent_email(i), scorecard_name(i),
                assessment_name(i), ""]
        for j, v in enumerate(vals):
            c = ws.cell(r, j + 1, v)
            c.border = BOX
            if j in (1, 2, 5):
                c.fill = FILL_INPUT
            if j in (3, 4):
                c.fill, c.font = FILL_GREY, F_NOTE
    ws.cell(AGENT_ROW0 + NUM_AGENTS + 1, 1,
            "Scorecard/Assessment Tab columns are maintained by the script — don't edit them by hand.").font = F_NOTE
    for col, w in zip("ABCDEF", (6, 30, 34, 24, 24, 24)):
        ws.column_dimensions[col].width = w
    ws.sheet_properties.tabColor = "6B7280"
    # Name in column B is used by other tabs; the label cells above sit in A:B for rows 4-11,
    # so agent names start at row 14 and never collide.


def build_dashboard(ws):
    cols = ["Agent", "Assessment Done", "Week Status", "RE Hours", "Prospecting Hrs", "Calls / Outreach",
            "DM Conversations", "Meetings", "New Opps", "New DB Contacts", "Follow-Ups Open", "MC Convos",
            "Pipeline GCI (latest)", "Q4 GCI Closed", "Q4 GCI Goal", "Q4 GCI %", "Q4 Closings", "Q4 Closing Goal",
            "Weekly Consistency", "Avg Calls / Wk", "Last Updated", "Win of the Week", "Biggest Obstacle"]
    title(ws, "Team Dashboard — Agent Score Card", len(cols),
          "Owner view. Change the “Week of” date to look at any week (defaults to the current week). "
          "Everything on this tab is calculated — agents update their own Scorecard tabs.")
    ws.cell(3, 1, "Week of:").font = F_BB
    wk = ws.cell(3, 2, "=TODAY()-WEEKDAY(TODAY(),3)")
    wk.number_format, wk.fill, wk.border, wk.font = DATE, FILL_INPUT, BOX, F_BB
    ws.cell(3, 3, "← type any Monday to review a past week").font = F_NOTE
    header_row(ws, 5, cols)
    ws.row_dimensions[5].height = 40
    last = FIRST_DATA + len(weeks()) - 1
    for i in range(1, NUM_AGENTS + 1):
        r = 5 + i
        sc, asm = q(scorecard_name(i)), q(assessment_name(i))
        rng = lambda name: f"{sc}!${COL[name]}${FIRST_DATA}:${COL[name]}${last}"  # noqa: E731
        dates = f"{sc}!$A${FIRST_DATA}:$A${last}"
        pick = lambda name: f'=IFERROR(INDEX({rng(name)},MATCH($B$3,{dates},0)),"")'  # noqa: E731
        name_ref = f"{q(SETTINGS)}!B{AGENT_ROW0 + i - 1}"
        f = [
            f'=IF({name_ref}="","Agent {i}",{name_ref})',
            f"={asm}!C3",
            f'=IFERROR(INDEX({rng("Status")},MATCH($B$3,{dates},0)),"")',
            pick("Total RE Hours"), pick("Prospecting Hours"), pick("Calls / Outreach Attempts"),
            pick("Decision-Maker Conversations"), pick("Prospect / Client Meetings"),
            pick("New Opportunities Created"), pick("New Database Contacts"), pick("Follow-Ups Outstanding"),
            pick("Market Center Agent Conversations"),
            f"={sc}!K10",
            f"={sc}!D6", f"={sc}!C6", f"={sc}!E6", f"={sc}!D8", f"={sc}!C8",
            f"={sc}!K7", f"={sc}!K8",
            f'=IF(MAX({rng("Last Updated (auto)")})=0,"",MAX({rng("Last Updated (auto)")}))',
            pick("Win of the Week"), pick("Biggest Obstacle"),
        ]
        for j, v in enumerate(f):
            c = ws.cell(r, j + 1, v)
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=True) if j >= 21 else CENTER
            if j == 0:
                c.font, c.alignment = F_BB, Alignment(vertical="center", indent=1)
            if i % 2 == 0:
                c.fill = FILL_LIGHT
        for j in (1, 15, 18):
            ws.cell(r, j + 1).number_format = PCT
        for j in (12, 13, 14):
            ws.cell(r, j + 1).number_format = MONEY
        ws.cell(r, 20).number_format = "0.0"
        ws.cell(r, 21).number_format = "mmm d, h:mm am/pm"
        ws.row_dimensions[r].height = 36
    tr = 6 + NUM_AGENTS
    ws.cell(tr, 1, "TEAM TOTAL").font = F_BB
    for j in range(4, 16):
        col = get_column_letter(j)
        if j in (16,):
            continue
        c = ws.cell(tr, j, f"=SUM({col}6:{col}{tr - 1})")
        c.font, c.fill, c.border, c.alignment = F_BB, FILL_GOLD, BOX, CENTER
        if j in (13, 14, 15):
            c.number_format = MONEY
    ws.cell(tr, 1).fill = FILL_GOLD
    for j in (17, 18):
        col = get_column_letter(j)
        c = ws.cell(tr, j, f"=SUM({col}6:{col}{tr - 1})")
        c.font, c.fill, c.border, c.alignment = F_BB, FILL_GOLD, BOX, CENTER
    status_rng = f"C6:C{tr - 1}"
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"⚠ Missing"'],
                                  fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color="991B1B", bold=True)))
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"✓ Updated"'],
                                  fill=PatternFill("solid", fgColor="D1FAE5"), font=Font(color="065F46")))
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Due Friday"'],
                                  fill=PatternFill("solid", fgColor="FEF3C7")))
    widths = [22, 11, 12, 9, 10, 10, 11, 9, 9, 10, 10, 9, 13, 12, 12, 9, 9, 10, 11, 10, 16, 34, 34]
    for j, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(j + 1)].width = w
    ws.freeze_panes = "B6"
    ws.sheet_properties.tabColor = "059669"


def build_review(ws):
    title(ws, "Praedium Development Review — Owner Only", NUM_AGENTS + 1,
          "Complete after meeting with each agent. Scores 1–5. This file is separate from the shared Score Card "
          "so agents never see it — do not share it.")
    header_row(ws, 4, ["Development Area"] + [agent_label(i) for i in range(1, NUM_AGENTS + 1)])
    areas = ["Available Time / Commitment", "Daily Structure", "Prospecting Consistency", "Pipeline Development",
             "Sphere / Relationship Development", "Market Center Development", "Market / Product Focus",
             "CRE Knowledge", "Organization / Follow-Through", "Coachability / Engagement"]
    r = 5
    ws.cell(r, 1, "Review meeting date").font = F_BB
    for i in range(NUM_AGENTS):
        c = ws.cell(r, i + 2)
        c.fill, c.border, c.number_format, c.alignment = FILL_INPUT, BOX, DATE, CENTER
    r += 1
    s = r
    for a in areas:
        ws.cell(r, 1, a).font = F_B
        ws.cell(r, 1).border = BOX
        for i in range(NUM_AGENTS):
            c = ws.cell(r, i + 2)
            c.fill, c.border, c.alignment = FILL_INPUT, BOX, CENTER
        r += 1
    list_dv(ws, ["1", "2", "3", "4", "5"], f"B{s}:{get_column_letter(NUM_AGENTS + 1)}{r - 1}")
    ws.cell(r, 1, "AVERAGE SCORE").font = F_BB
    for i in range(NUM_AGENTS):
        col = get_column_letter(i + 2)
        c = ws.cell(r, i + 2, f'=IFERROR(AVERAGE({col}{s}:{col}{r - 1}),"")')
        c.font, c.fill, c.border, c.number_format, c.alignment = F_BB, FILL_GOLD, BOX, "0.0", CENTER
    ws.cell(r, 1).fill = FILL_GOLD
    ws.conditional_formatting.add(f"B{s}:{get_column_letter(NUM_AGENTS + 1)}{r}", CellIsRule(
        operator="lessThanOrEqual", formula=["2"], fill=PatternFill("solid", fgColor="FEE2E2")))
    ws.conditional_formatting.add(f"B{s}:{get_column_letter(NUM_AGENTS + 1)}{r}", CellIsRule(
        operator="greaterThanOrEqual", formula=["4"], fill=PatternFill("solid", fgColor="D1FAE5")))
    r += 2
    for lab in ["Agent's strongest area", "Primary obstacle to growth", "Highest-priority development area",
                "What Praedium should provide", "30-Day Objective", "90-Day Objective"]:
        ws.cell(r, 1, lab).font = F_BB
        ws.cell(r, 1).alignment = WRAP
        for i in range(NUM_AGENTS):
            c = ws.cell(r, i + 2)
            c.fill, c.border, c.alignment = FILL_INPUT, BOX, WRAP
        ws.row_dimensions[r].height = 60
        r += 1
    ws.column_dimensions["A"].width = 34
    for i in range(NUM_AGENTS):
        ws.column_dimensions[get_column_letter(i + 2)].width = 24
    ws.freeze_panes = "B5"
    ws.sheet_properties.tabColor = "991B1B"


def build_start(ws):
    title(ws, "THE PRAEDIUM GROUP — Agent Score Card", 3)
    lines = [
        ("HOW THIS WORKS", None),
        ("1.", "You have two tabs with your name on them: your Weekly Scorecard and your one-time Assessment."),
        ("2.", "Complete your Assessment first (target: before our next meeting). Estimates are fine — there are no "
               "“right” answers. Your Q4 goals in Section 13 drive the progress tracker on your Scorecard."),
        ("3.", "Every week, fill in your row on your Scorecard by Friday end of day. It takes about 5 minutes."),
        ("4.", "You'll get an email reminder each Friday with a direct link to your tab. If last week is still "
               "blank on Monday, you'll get a gentle nudge."),
        ("5.", "Only yellow cells are for you. Gray/blue cells calculate automatically. You can only edit your own tabs."),
        ("", None),
        ("WHAT TO TRACK EACH WEEK", None),
        ("Calls / Outreach", "Every dial, email, text, or door knock aimed at a prospect (owner, tenant, investor)."),
        ("DM Conversations", "Real conversations with a decision maker — not voicemails or gatekeepers."),
        ("Meetings", "Discovery, listing, buyer/tenant, or prospect meetings held (in person or video)."),
        ("New Opportunities", "A qualified prospect with a real need, timeline, and next step."),
        ("MC Conversations", "Meaningful conversations with agents in your KW Market Center about commercial."),
        ("Sphere Touches", "Intentional contacts with attorneys, CPAs, lenders, developers, past clients, etc."),
        ("Pipeline GCI", "Your best estimate of total potential GCI in your active pipeline at week's end."),
        ("", None),
        ("QUESTIONS?", "Reach out any time — this is a tool to help you grow, not a report card."),
    ]
    r = 3
    for a, b in lines:
        if b is None and a:
            ws.cell(r, 1, a).font = F_SEC
        else:
            ws.cell(r, 1, a).font = F_BB
            c = ws.cell(r, 2, b)
            c.font, c.alignment = F_B, WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            if b and len(b) > 90:
                ws.row_dimensions[r].height = 30
        r += 1
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 40
    ws.sheet_properties.tabColor = GOLD


def main():
    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    dash = wb.create_sheet("Dashboard")
    settings = wb.create_sheet(SETTINGS)
    build_start(start)
    build_settings(settings)
    for i in range(1, NUM_AGENTS + 1):
        sc = wb.create_sheet(scorecard_name(i))
        asm = wb.create_sheet(assessment_name(i))
        keys = build_assessment(asm, i)
        build_scorecard(sc, i, keys)
    build_dashboard(dash)

    private = Workbook()
    build_review(private.active)
    private.active.title = "Praedium Review"

    for book, fname in ((wb, "Praedium_Agent_Score_Card.xlsx"), (private, "Praedium_Development_Review_PRIVATE.xlsx")):
        for ws in book.worksheets:
            ws.sheet_view.showGridLines = False
        out = Path(__file__).with_name(fname)
        book.save(out)
        print(f"wrote {out}")


if __name__ == "__main__":
    main()
